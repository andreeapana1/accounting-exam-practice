"""More Excel questions on the same course files (second batch). Answers are computed from the data."""
import csv
import re
from collections import Counter, defaultdict
from datetime import datetime

import openpyxl

from common import Bank, money, num
from excel_questions import ETL, INTEGRATECO, PWC, _case_rows, _parse_case3, _read_pipe, jobcode


def unique_top(counter, what):
    """Most common item, refusing ties so every answer is unambiguous."""
    (a, n), *rest = counter.most_common(2) + [(None, -1)]
    assert rest[0][1] < n, f"tie for {what}: {counter.most_common(3)}"
    return a, n


def top_all(counter):
    """Highest count and every key that reaches it (for questions where ties are part of the answer)."""
    n = max(counter.values())
    return n, sorted(k for k, v in counter.items() if v == n)


def nyc():
    b = Bank("nyc", "Excel")
    with open(PWC / "NYC_Film_Permits.txt", newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    t = lambda s: datetime.strptime(s, "%m/%d/%Y %H:%M")
    for r in rows:
        r["s"], r["en"] = t(r["StartDateTime"]), t(r["EnteredOn"])
    hour, hn = unique_top(Counter(r["s"].hour for r in rows), "start hour")
    weekend_entry = sum(r["en"].weekday() >= 5 for r in rows)
    boards = [len([x for x in r["CommunityBoard(s)"].split(",") if x.strip()]) for r in rows]
    zipc, zn = unique_top(Counter(r["ZipCode"] for r in rows if r["ZipCode"] != "N/A"), "zip")
    segments5 = sum(len([x for x in r["ParkingHeld"].split(",") if x.strip()]) >= 5 for r in rows)
    single = Counter(r["PolicePrecinct(s)"].strip() for r in rows if "," not in r["PolicePrecinct(s)"])
    prec, pn = unique_top(single, "precinct")
    on_hour = sum(r["s"].minute == 0 for r in rows)

    b.add("At which hour of the day do the most permits start?", f"{hour}:00 ({num(hn)} permits)",
          '=HOUR(C2) in a helper column, then a PivotTable count (or =MODE.SNGL on the helper column).', 2,
          ["dates", "PivotTable"])
    b.add("How many permits were entered (EnteredOn) on a Saturday or Sunday?", num(weekend_entry),
          "=WEEKDAY(E2,2)>=6 in a helper column, then count TRUE.", 2, ["WEEKDAY", "COUNTIF"])
    b.add("On average, how many community boards does a permit involve? (2 decimals)", f"{sum(boards) / len(boards):.2f}",
          'Boards per permit =LEN(I2)-LEN(SUBSTITUTE(I2,",",""))+1, then =AVERAGE() of that column.', 2.5,
          ["LEN", "SUBSTITUTE", "AVG"])
    b.add("Which zip code appears on the most permits?", f"{zipc} ({num(zn)} permits)",
          'PivotTable with ZipCode in Rows and a count, sorted descending (exclude "N/A").', 1.5, ["PivotTable"])
    b.add("How many permits hold parking on five or more street segments (comma-separated in ParkingHeld)?",
          num(segments5), 'Segments =LEN(G2)-LEN(SUBSTITUTE(G2,",",""))+1; then =COUNTIF(P:P,">=5").', 2.5,
          ["LEN", "SUBSTITUTE", "COUNTIF"])
    b.add("Among permits that list a single police precinct, which precinct appears most often?",
          f"Precinct {prec} ({num(pn)} permits)", 'Filter out values containing a comma, then a PivotTable count.', 2.5,
          ["filtering", "PivotTable", "wildcards"])
    b.add("How many permits start exactly on the hour (minute 00)?", num(on_hour),
          "=MINUTE(C2)=0 in a helper column, then count TRUE.", 1.5, ["dates", "COUNTIF"])
    cross_year = sum(t(r["EndDateTime"]).year != r["s"].year for r in rows)
    b.add("How many permits end in a different calendar year from the one they start in?", cross_year,
          "=YEAR(D2)<>YEAR(C2) in a helper column, then count TRUE (these shoots run over New Year).", 2,
          ["YEAR", "dates"])
    multi = defaultdict(lambda: [0, 0])
    for r in rows:
        multi[r["Borough"]][0] += "," in r["CommunityBoard(s)"]
        multi[r["Borough"]][1] += 1
    share = {k: a / n for k, (a, n) in multi.items()}
    top_b = max(share.items(), key=lambda kv: kv[1])
    b.add("Which borough has the highest share of permits that need more than one community board? (1 decimal)",
          f"{top_b[0]} ({top_b[1] * 100:.1f}%)",
          'Flag =ISNUMBER(SEARCH(",",I2)), then a PivotTable by Borough showing the average of the flag (TRUE = 1).', 3,
          ["PivotTable", "percentages", "wildcards"])
    return b.items


def _etl1_clean():
    with open(ETL / "case 1" / "Analytics_mindset_case_studies_ETL_Case1.csv", newline="", encoding="utf-8-sig") as f:
        raw = list(csv.reader(f))
    header, rows = raw[0], [r for r in raw[1:] if any(c.strip() for c in r)]
    fixed = []
    for r in rows:
        if len(r) > len(header):  # unquoted comma in "Apple, Inc."
            r = [r[0], r[1] + "," + r[2]] + r[3:]
        fixed.append(r)
    return header, fixed


def _money_value(s):
    s = s.strip().replace('"', "")
    mult = 1
    if s.endswith(" M"):
        mult, s = 1_000_000, s[:-2]
    elif s.endswith(" K"):
        mult, s = 1_000, s[:-2]
    neg = s.startswith("(")
    s = s.strip("()").replace(",", "").strip()
    if s in ("", "-"):
        return None
    v = float(s) * mult
    return -v if neg else v


def etl1():
    b = Bank("etl1", "Excel")
    header, rows = _etl1_clean()
    names = {r[1].strip() for r in rows}
    cash18 = {r[1].strip(): _money_value(r[3]) for r in rows if r[2] == "2018"}
    top = max(cash18.items(), key=lambda kv: kv[1])
    years = Counter(r[2] for r in rows)
    b.add("How many different companies (by Name) are in the balance sheet file?", len(names),
          "=COUNTA(UNIQUE(B2:B23)) after fixing the Apple rows; tickers cannot be used because some are blank or contain spaces.",
          1.5, ["UNIQUE", "unique identifiers"])
    b.add("After converting K and M suffixes to full numbers, which company had the most Cash in 2018?",
          f"{top[0]} ({num(top[1])})", "Convert Cash with =IF(RIGHT(D2,1)=\"M\",VALUE(LEFT(D2,LEN(D2)-2))*10^6,IF(RIGHT(D2,1)=\"K\","
          "VALUE(LEFT(D2,LEN(D2)-2))*10^3,VALUE(D2))), filter Year = 2018, then MAX.", 3,
          ["units of measurement", "IF", "VALUE"])
    b.add("Which year has more rows than the other years, and why?",
          f"{years.most_common(1)[0][0]} ({years.most_common(1)[0][1]} rows) — it contains the duplicated row",
          "PivotTable count by Year; compare with the number of companies to spot the duplicate.", 2,
          ["duplicates", "PivotTable"])
    return b.items


def etl2():
    b = Bank("etl2", "Excel")
    rows = _case_rows(ETL / "case 2" / "Analytics_mindset_case_studies_ETL_Case2_Excel.xlsx", "Case 2")
    parsed = []
    for code, first, last in rows:
        m = re.fullmatch(r"([A-Z]{3})(\d+)-(\d+)([WM])", code)
        parsed.append(dict(loc=m[1], emp=int(m[2]), plant=int(m[3]), pay=m[4], first=first, last=last))
    plants = defaultdict(set)
    for p in parsed:
        plants[p["loc"]].add(p["plant"])
    most_n, most_who = top_all(Counter({k: len(v) for k, v in plants.items()}))
    per_plant = Counter((p["loc"], p["plant"]) for p in parsed)
    lonely = sum(n == 1 for n in per_plant.values())
    share = {loc: sum(p["pay"] == "M" for p in parsed if p["loc"] == loc) / sum(p["loc"] == loc for p in parsed)
             for loc in plants}
    top_share = max(share.items(), key=lambda kv: kv[1])
    can = min((p for p in parsed if p["loc"] == "CAN"), key=lambda p: p["emp"])

    b.add("What is the largest number of different plants any country operates, and which countries reach it?",
          f"{most_n} plants: {', '.join(most_who)}",
          'Build LocationPlantID, then count distinct plants per country: =COUNTA(UNIQUE(FILTER(F:F,D:D="ARG"))) or a '
          "PivotTable with distinct count.", 3, ["UNIQUE", "PivotTable", "concatenation"])
    b.add("How many plants (LocationPlantID) have exactly one employee?", lonely,
          "PivotTable on LocationPlantID with a count, then count the rows equal to 1 (or =COUNTIF on the counts).", 3,
          ["PivotTable", "COUNTIF"])
    b.add("Which country has the highest share of monthly-paid employees? (1 decimal)",
          f"{top_share[0]} ({top_share[1] * 100:.1f}%)",
          'Per country: =COUNTIFS(D:D,"USA",G:G,"M")/COUNTIF(D:D,"USA"), or a PivotTable showing values as % of row total.',
          3, ["COUNTIFS", "percentages", "PivotTable"])
    b.add("Who has the lowest EmpID in Canada, and what is it?", f"{can['first']} {can['last']} ({can['emp']})",
          '=MINIFS(E:E,D:D,"CAN") then XLOOKUP the name.', 2, ["MINIFS", "XLOOKUP"])
    return b.items


def etl3():
    b = Bank("etl3", "Excel")
    rows = _case_rows(ETL / "case 3" / "Analytics_mindset_case_studies_ETL_Case3_Excel.xlsx", "Case 3 data")
    parsed = [dict(_parse_case3(c), first=f, last=l) for c, f, l in rows]
    plants = defaultdict(set)
    for p in parsed:
        plants[p["loc"]].add(p["plant"])
    most_n, most_who = top_all(Counter({k: len(v) for k, v in plants.items()}))
    lonely = sum(n == 1 for n in Counter((p["loc"], p["plant"]) for p in parsed).values())
    aus_weekly = sum(p["loc"] == "AUS" and p["pay"] == "W" for p in parsed)
    top = max(parsed, key=lambda p: p["emp"])
    first_dupes = sum(n > 1 for n in Counter(p["first"] for p in parsed).values())

    b.add("In Case 3, what is the largest number of different plants any country operates, and which countries reach it?",
          f"{most_n} plants: {', '.join(most_who)}",
          "Extract Location and PlantID, then a distinct count of PlantID per Location (PivotTable or UNIQUE + FILTER).", 3.5,
          ["LEFT", "MID", "FIND", "UNIQUE"])
    b.add("How many plants (LocationPlantID) have exactly one employee in Case 3?", lonely,
          "Build LocationPlantID from the extracted fields, then a PivotTable count and count the 1s.", 3.5,
          ["concatenation", "PivotTable"])
    b.add("How many employees in Australia are paid weekly (W or We)?", aus_weekly,
          'Normalise pay period with =IF(LEFT(G2,1)="W","W","M"), then =COUNTIFS(D:D,"AUS",H:H,"W").', 2.5,
          ["IF", "COUNTIFS"])
    b.add("Which employee has the highest EmpID, and what is it?", f"{top['first']} {top['last']} ({top['emp']})",
          'Extract EmpID with MID and FIND, convert with VALUE, then =MAX() and XLOOKUP the name.', 3,
          ["MID", "FIND", "VALUE", "XLOOKUP"])
    b.add("How many first names are shared by more than one employee?", first_dupes,
          "=COUNTIF(B:B,B2)>1 in a helper column, then count the distinct names flagged TRUE (or a PivotTable).", 2,
          ["COUNTIF", "unique identifiers"])
    return b.items


def etl4():
    b = Bank("etl4", "Excel")
    wb = openpyxl.load_workbook(ETL / "case 4" / "Analytics_mindset_case_studies_ETL_Case4_Excel.xlsx", data_only=True)
    je = [r for r in wb["JELineItems"].iter_rows(min_row=2, values_only=True) if r[0]]
    coa = {r[1]: (r[0].strip(), r[2], r[3]) for r in wb["ChartOfAccounts"].iter_rows(min_row=2, values_only=True) if r[1]}
    f = lambda v: float(v) if v not in (None, "") else 0.0
    lines = Counter(r[0] for r in je)
    many = sum(n > 2 for n in lines.values())
    day, dn = unique_top(Counter(r[5].date() for r in {r[0]: r for r in je}.values()), "day")
    big = min(je, key=lambda r: f(r[8]))
    assets = sum(1 for r in je if coa[r[1]][1] == "Assets")
    preparers = len({(r[3], r[2]) for r in je})
    payroll_desc = len({r[0] for r in je if "PAYROLL" in (r[6] or "").upper()})

    src = {r[0]: r[1] for r in wb["Source"].iter_rows(min_row=2, values_only=True) if r[0]}
    pcard_jes = len({r[0] for r in je if src.get(r[4]) == "PURCHASE CARD"})
    b.add("How many journal entries (distinct JENumbers) come from the PURCHASE CARD source?", pcard_jes,
          "Look up Source with XLOOKUP on SourceID, filter on PURCHASE CARD, then count distinct JENumbers.", 2.5,
          ["XLOOKUP", "DISTINCT"])
    b.add("How many journal entries have more than two lines?", many,
          "=COUNTIF(A:A,A2) gives the lines per entry; count distinct JENumbers where it is above 2 (or a PivotTable).", 2.5,
          ["COUNTIF", "PivotTable"])
    b.add("On which EffectiveDate were the most journal entries (distinct JENumbers) posted?",
          f"{day.month}/{day.day}/{day.year} ({dn} journal entries)",
          "PivotTable with EffectiveDate in rows and a distinct count of JENumber (add to Data Model), sorted descending.", 3,
          ["PivotTable", "dates", "DISTINCT"])
    b.add("What is the largest single credit line, and to which GL account was it posted?",
          f"{money(f(big[8]))} to {coa[big[1]][0]} ({big[1]})",
          "=MIN(I:I) (credits are negative), then look up GLAccountNumber and its name with XLOOKUP.", 2.5,
          ["MIN", "XLOOKUP"])
    b.add("How many line items are posted to accounts whose AccountType is Assets?", assets,
          'Look up AccountType for each line with XLOOKUP on GLAccountNumber, then =COUNTIF(N:N,"Assets").', 2.5,
          ["XLOOKUP", "COUNTIF"])
    b.add("How many different preparers appear in the journal? (Remember PreparerID is unique only within a business unit.)",
          preparers, "Combine =D2&\"-\"&C2 (PreparerID and BusinessUnitID), then =COUNTA(UNIQUE(...)).", 2.5,
          ["unique identifiers", "concatenation", "UNIQUE"])
    b.add("How many journal entries have a description containing PAYROLL?", payroll_desc,
          'Flag lines with =ISNUMBER(SEARCH("PAYROLL",G2)), then count distinct JENumbers among the flagged lines.', 2.5,
          ["wildcards", "DISTINCT"])
    return b.items


def integrateco():
    b = Bank("integrateco", "Excel")
    p15 = _read_pipe("Analytics_mindset_case_studies_IntegrateCo_2015payroll.csv")
    p16 = _read_pipe("Analytics_mindset_case_studies_IntegrateCo_2016payroll.csv")
    jobs = {r["Job_Code"]: r["Job_Description"] for r in _read_pipe("Analytics_mindset_case_studies_IntegrateCo_jobcodes.csv")}
    f = lambda v: float(v or 0)
    paid_no_hours = sum(f(r["Reg_Pay"]) > 0 and f(r["Reg_Hrs"]) == 0 for r in p15)
    ot_share = sum(f(r["OT_Pay"]) for r in p15) / sum(f(r["Reg_Pay"]) + f(r["OT_Pay"]) for r in p15)
    g15, g16 = defaultdict(float), defaultdict(float)
    for r in p15:
        g15[jobcode(r["Job"])] += f(r["Reg_Pay"]) + f(r["OT_Pay"])
    for r in p16:
        g16[jobcode(r["Project"])] += f(r["Gross_Wages"])
    inc = max(((j, g16[j] - g15[j]) for j in jobs), key=lambda kv: kv[1])
    ids = defaultdict(set)
    for r in p16:
        if r["Employee_Number"].strip():
            ids[r["Employee_Number"]].add(r["Employee_ID"])
    multi_id = sum(len(v) > 1 for v in ids.values())
    taxes15 = sum(f(r["Futa"]) + f(r["Suta"]) for r in p15)
    taxes16 = sum(f(r["FUTA"]) + f(r["SUTA"]) for r in p16)
    rate005 = (sum(f(r["Reg_Pay"]) for r in p15 if jobcode(r["Job"]) == "005")
               / sum(f(r["Reg_Hrs"]) for r in p15 if jobcode(r["Job"]) == "005"))

    b.add("Data check: how many 2015 rows have regular pay above zero but zero regular hours?", paid_no_hours,
          '=COUNTIFS(K:K,">0",I:I,0) on the 2015 data.', 1.5, ["COUNTIFS", "data quality"])
    b.add("What share of 2015 gross pay was overtime pay? (2 decimals)", f"{ot_share * 100:.2f}%",
          "=SUM(L:L)/(SUM(K:K)+SUM(L:L)) on the 2015 data.", 1.5, ["SUM", "percentages"])
    b.add("Which job code had the largest increase in gross wages from 2015 to 2016, and by how much?",
          f"{inc[0]} {jobs[inc[0]]} (+{money(inc[1])})",
          "Extract the job code in both years, total gross wages per code with SUMIF (2015: Reg_Pay + OT_Pay), then "
          "subtract and sort.", 4, ["MID", "SUMIF", "comparability"])
    b.add("Data check: how many 2016 Employee_Numbers are linked to more than one Employee_ID?",
          f"{multi_id}" + (" — the new system's Employee_ID is consistent" if multi_id == 0 else ""),
          "PivotTable: Employee_Number in rows, distinct count of Employee_ID; count rows above 1.", 3,
          ["unique identifiers", "PivotTable"])
    b.add("By how much did employer unemployment taxes (FUTA + SUTA) change from 2015 to 2016?",
          f"{money(taxes16 - taxes15)} (2015: {money(taxes15)}; 2016: {money(taxes16)})",
          "Sum Futa and Suta in 2015 and FUTA and SUTA in 2016, then subtract; note the different column names.", 2,
          ["SUM", "comparability"])
    run = defaultdict(float)
    for r in p15:
        run[r["Pay_Date"]] += f(r["Reg_Pay"]) + f(r["OT_Pay"])
    top_run = max(run.items(), key=lambda kv: kv[1])
    b.add("Which 2015 payroll run (Pay_Date) had the highest total gross pay, and how much?",
          f"{top_run[0]} ({money(top_run[1])})",
          "Gross pay =K2+L2 per row, then a PivotTable with Pay_Date in rows and the sum, sorted descending.", 2.5,
          ["PivotTable", "dates"])
    ot_emps = len({r["Emp_No"] for r in p15 if f(r["OT_Hrs"]) > 0})
    b.add("How many different employees worked any overtime in 2015?", ot_emps,
          '=COUNTA(UNIQUE(FILTER(C2:C4982,J2:J4982>0))) — count employees, not rows.', 2,
          ["UNIQUE", "filtering", "COUNT"])
    b.add("What was the average regular pay per regular hour for job code 005 (Installation) in 2015? (2 decimals)",
          f"{rate005:.2f}", 'Extract the job code, then =SUMIF(job,"005",Reg_Pay)/SUMIF(job,"005",Reg_Hrs) — '
          "not the average of row-level rates.", 3, ["SUMIF", "MID", "aggregation"])
    return b.items


def all_questions():
    return nyc() + etl1() + etl2() + etl3() + etl4() + integrateco()
