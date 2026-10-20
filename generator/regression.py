"""Data interpretation (regression) exams.

Two exams use regressions actually estimated on the restaurant case data; the rest are business scenarios
from the course settings with constructed coefficients. Every exam has one model and four decision-oriented
questions: a planning prediction, the size of a group difference, which variable to drop, and whether a
prediction can be trusted.
"""
import random
from collections import defaultdict

import numpy as np
import openpyxl
from scipy import stats

from common import EXCEL, num

RESTAURANT = EXCEL / "4. Regression Assignment" / "Restaurant" / "data.xlsm"
SUB = "₀₁₂₃"


def fmt(x):
    """Coefficient display: no trailing zeros, at most 3 decimals."""
    s = f"{x:,.3f}".rstrip("0").rstrip(".")
    return "0" if s == "-0" else s


def signed(c):
    """' + 5' or ' − 5' for writing out an equation."""
    return f" − {fmt(abs(c))}" if c < 0 else f" + {fmt(c)}"


def tidy(c):
    """Round an estimated coefficient for display; predictions use the rounded value."""
    return round(float(c), 3 if abs(c) < 1 else 2)


def ols(y, cols):
    X = np.column_stack([np.ones(len(y))] + cols)
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)
    resid = y - X @ beta
    dof = len(y) - X.shape[1]
    se = np.sqrt(np.diag(resid @ resid / dof * np.linalg.inv(X.T @ X)))
    p = 2 * stats.t.sf(np.abs(beta / se), dof)
    r2 = 1 - resid @ resid / ((y - y.mean()) @ (y - y.mean()))
    return beta, p, r2, len(y)


# ------------------------------------------------------------------ exam text
def build_exam(eid, title, context, y, x, d, coefs, pvals, interaction, predict_at, gap_at, outside, source):
    """y/x: dicts with name, unit; d: dict with name, ref, comp (labels); coefs/pvals: lists."""
    letter = "d" if interaction else "b"
    par = [f"{letter}{SUB[i]}" for i in range(len(coefs))]
    terms = ["Intercept", x["name"], d["name"]] + ([f'{x["name"]}×{d["name"]}'] if interaction else [])
    eq = f'Predicted {y["name"].lower()} = {par[0]} + {par[1]}[{x["name"]}] + {par[2]}[{d["name"]}]'
    if interaction:
        eq += f' + {par[3]}([{x["name"]}]×[{d["name"]}])'

    coefs = [tidy(c) for c in coefs]
    b0, b1, b2 = coefs[:3]
    b3 = coefs[3] if interaction else 0
    unit = y["unit"]
    # How to say a value of x in a sentence. Some variables are measured in hundreds or thousands
    # (x = 46 means 4,600 invoices), so sentences use the real amount and answers explain the scaling.
    scale = x.get("scale", 1)
    say = lambda v: x["phrase"].format(v=num(v * scale))
    joined = lambda group: f"{group}{' and ' if ' with ' in group else ' with '}"
    lo, hi = x["range"]
    per, _, base = x["unit"].partition(" ")  # "hundred invoices" -> "hundred", "invoices"
    scale_note = lambda v: (f" {x['name']} is measured in {per}s of {base}, so {num(v * scale)} means "
                            f"{x['name']} = {num(v)}." if scale != 1 else "")
    span = (f"{x['name'].lower()} from {num(lo * scale)} to {num(hi * scale)} "
            f"{base if scale != 1 else x['unit']}")
    who = context["stakeholder"]
    # "Noah, IntegrateCo's CFO" needs a closing comma when it is the subject of a sentence.
    subj = who + ("," if "," in who else "")
    Who = subj[0].upper() + subj[1:]
    lo, hi = x["range"]
    unit_words = {"USD": ["dollar", "usd"], "EUR": ["euro", "eur"], "DKK": ["kroner", "dkk"]}
    unit_terms = next((v for k, v in unit_words.items() if k in unit), [unit.split()[0].lower()])
    ref_terms = ["reference", "compared with", "compared to", "than "] + [w for w in d["ref"].lower().split() if len(w) > 5][:2]

    # 1. Planning prediction
    xv, dv = predict_at
    yhat = b0 + b1 * xv + b2 * dv + b3 * xv * dv
    group = d["comp"] if dv else d["ref"]
    work = f"{fmt(b0)}{signed(b1)} × {num(xv)}{signed(b2)} × {dv}"
    if interaction:
        work += f"{signed(b3)} × {num(xv)} × {dv}"
    dropped = "" if dv else f" The {d['name'].lower()} term{'s' if interaction else ''} drop{'' if interaction else 's'} out because {d['name']} = 0."
    q_plan = {
        "kind": "P",
        "prompt": f"{Who} is planning for {joined(group)}{say(xv)}. What {y['name'].lower()} should they plan for?",
        "answer": f"{scale_note(xv).strip()} Predicted {y['name'].lower()} = {work} = {fmt(yhat)} {unit}.{dropped}".strip(),
        "keypoints": [
            {"label": f"Gave the prediction ({fmt(yhat)})", "number": yhat, "tol": 0.01,
             "probe": "So what is the figure? Give me the actual number you would plan for."},
            {"label": f"Stated the unit ({unit})", "terms": unit_terms, "probe": "In what unit is that number?"},
            {"label": "Showed how the values go into the equation", "terms": ["plus", "times", "multipl", "+", "×"],
             "probe": "How did you get there? Which values go into the equation?"},
        ],
    }

    # 2. Size of the group difference
    gx = gap_at
    gap = b2 + b3 * gx
    if interaction:
        gap_ans = (f"With the interaction, the gap depends on {x['name'].lower()}: gap = {par[2]} + {par[3]} × "
                   f"{x['name']} = {fmt(b2)}{signed(b3)} × {num(gx)} = {fmt(gap)} {unit}.{scale_note(gx)} "
                   f"With {say(gx)}, {d['comp']} are predicted to have {fmt(abs(gap))} {unit} "
                   f"{'more' if gap >= 0 else 'less'} than {d['ref']}.")
    else:
        gap_ans = (f"Without an interaction the gap is the same at every level of {x['name'].lower()}: it is {par[2]} = "
                   f"{fmt(b2)} {unit}. Holding {x['name'].lower()} constant, {d['comp']} are predicted to have "
                   f"{fmt(abs(b2))} {unit} {'more' if b2 >= 0 else 'less'} than {d['ref']} (the reference group).")
    q_gap = {
        "kind": "G",
        "prompt": f"With {say(gx)}, how large is the predicted difference in {y['name'].lower()} between "
                  f"{d['comp']} and {d['ref']}? Which part of the equation gives you that number?",
        "answer": gap_ans,
        "keypoints": [
            {"label": f"Gave the size of the gap ({fmt(gap)})", "number": gap, "tol": 0.01,
             "probe": "So how big is the difference, as a number?"},
            {"label": f"Named the comparison group ({d['ref']})", "terms": ref_terms,
             "probe": "More or less compared with which group, exactly?"},
            {"label": "Pointed to the right coefficient(s)" + (" — including the interaction" if interaction else ""),
             "terms": ["interaction", "slope", "depends"] if interaction else ["coefficient", par[2], "dummy", "constant"],
             "probe": "Does that gap stay the same for every value, or does it change?" if interaction
             else "Which number in the regression table gives you that difference?"},
        ],
    }

    # 3. Which variable to drop
    named = list(zip(terms[1:], pvals[1:]))
    weakest = max(named, key=lambda tp: tp[1])
    if weakest[1] >= 0.05:
        drop_ans = (f"{weakest[0]} is the candidate to drop: its p-value ({weakest[1]:.3g}) is above 0.05, so we cannot rule "
                    f"out that its true effect is zero.")
        if interaction and weakest[0] != terms[3]:
            drop_ans += " Careful: if the interaction stays in the model, its main effects should normally stay too."
    else:
        drop_ans = ("None — every predictor has a p-value below 0.05, so each one adds explanatory power. Dropping one "
                    "would make the model simpler but worse.")
    drop_ans += (" Statistical significance only says an effect is unlikely to be zero; whether it matters for the business "
                 "depends on its size.")
    q_drop = {
        "kind": "S",
        "prompt": f"{Who} wants a simpler model and asks you to remove one variable. Based on the regression output, "
                  f"which one would you remove — or would you keep them all? Why?",
        "answer": drop_ans,
        "keypoints": [
            {"label": "Used the p-values / 5% level", "terms": ["0.05", "5%", "5 %", "five percent", "p-value", "p value"],
             "probe": "What in the regression output are you basing that on?"},
            {"label": f"Chose {weakest[0] if weakest[1] >= 0.05 else 'to keep all variables'}",
             "terms": [w.lower() for w in weakest[0].replace("×", " ").split() if len(w) > 3][:2] + (["interaction"] if "×" in weakest[0] else [])
             if weakest[1] >= 0.05 else ["keep", "all of them", "none"],
             "probe": "Look at the p-values again: which variable has the weakest evidence?" if weakest[1] >= 0.05
             else "Is there any variable you would remove, and why or why not?"},
            {"label": "Explained what a high/low p-value means", "terms": ["zero", "no effect", "chance", "rule out", "evidence"],
             "probe": "What does a p-value above 0.05 actually tell you about that variable?"},
        ],
    }

    # 4. Can the prediction be trusted?
    ox, odv = outside
    oyhat = b0 + b1 * ox + b2 * odv + b3 * ox * odv
    q_trust = {
        "kind": "R",
        "prompt": f"A colleague used the model to predict {y['name'].lower()} for {joined(d['comp'] if odv else d['ref'])}"
                  f"{say(ox)} and got {fmt(oyhat)} {unit}. Should {subj} rely on this number?",
        "answer": f"No, not without a strong warning. The model was estimated on data with {span}, so "
                  f"{say(ox)} lies outside the observed range. This is extrapolation: the linear relationship may not "
                  f"hold there. Use values inside the range, or collect data that covers this situation.",
        "keypoints": [
            {"label": "Said the value is outside the observed range", "terms": ["outside", "range", "beyond", "higher than", "above"],
             "probe": "What range of values was the model actually estimated on?"},
            {"label": "Called it extrapolation / unreliable", "probe": "What is the risk of using the model for a value like that?", "terms": ["extrapolat", "unreliable", "not reliable", "careful",
                                                                       "caution", "uncertain", "not confident", "risky"]},
            {"label": f"Mentioned the observed range ({num(lo * scale)}–{num(hi * scale)})", "number": hi * scale,
             "alt": [hi] if scale != 1 else [], "tol": 0.01, "probe": "Can you tell me the range the data covers?"},
        ],
    }
    return {
        "id": eid, "title": title, "intro": context["intro"], "source": source, "stakeholder": who,
        "y": y, "x": x, "d": d, "equation": eq,
        "table": [{"param": p, "term": t, "estimate": fmt(c), "p": f"{pv:.3f}".rstrip("0").rstrip(".") if pv >= 0.001 else "<0.001"}
                  for p, t, c, pv in zip(par, terms, coefs, pvals)],
        "questions": [q_plan, q_gap, q_drop, q_trust],
        "minutes": 5,
    }


# --------------------------------------------------------- restaurant (real)
def restaurant_data():
    wb = openpyxl.load_workbook(RESTAURANT, read_only=True, data_only=True)
    sales, sales364 = defaultdict(float), defaultdict(float)
    for r in wb["salesData"].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        sales[r[0].date()] += r[5] or 0
        sales364[r[0].date()] += r[11] or 0
    tmax = {r[0].date(): r[3] for r in wb["weatherData"].iter_rows(min_row=2, values_only=True) if r[0]}
    game = {r[0].date(): r[4] for r in wb["gameData"].iter_rows(min_row=2, values_only=True) if r[0]}
    days = sorted(d for d in sales if d in tmax)
    return {
        "y": np.array([sales[d] for d in days]),
        "y364": np.array([sales364[d] for d in days]),
        "tmax": np.array([tmax[d] for d in days]),
        "game": np.array([float(game.get(d, 0)) for d in days]),
    }


def restaurant_exams():
    r = restaurant_data()
    stake = {"stakeholder": "Bob, the restaurant owner",
             "intro": "Larry's Commissary point-of-sale data for 2018, merged with daily weather and university football "
                      "game days (the restaurant regression case). The regression was estimated on the case data."}
    y = {"name": "Daily sales", "unit": "USD"}
    game = {"name": "Game day", "ref": "days without a home-town football game", "comp": "game days"}
    t = {"name": "Max temperature", "unit": "°F", "unit_one": "degree Fahrenheit", "phrase": "a maximum temperature of {v} °F",
         "range": (round(float(r["tmax"].min())), round(float(r["tmax"].max())))}
    beta, p, r2, n = ols(r["y"], [r["tmax"], r["game"]])
    ex1 = build_exam("REG01", "Larry's Commissary: weather and game days", stake, y, t, game, beta, p, False,
                     (75, 1), 75, (110, 0), f"Estimated on the case data: n = {n} days, R² = {r2:.2f}")

    hi = float(r["y364"].max())
    s364 = {"name": "Sales one year ago", "unit": "USD", "unit_one": "dollar of sales on the same day last year",
            "phrase": "sales of {v} USD on the same day last year",
            "range": (round(float(r["y364"].min()), -2), round(hi, -2))}
    beta, p, r2, n = ols(r["y"], [r["y364"], r["game"], r["y364"] * r["game"]])
    ex2 = build_exam("REG02", "Larry's Commissary: last year's sales × game day", stake, y, s364, game, beta, p, True,
                     (4000, 1), 5000, (round(hi * 1.6, -2), 1), f"Estimated on the case data: n = {n} days, R² = {r2:.2f}")
    return [ex1, ex2]


# ------------------------------------------------------- scenario templates
SCENARIOS = [
    dict(title="Month-end close duration", phrase="{v} journal entries", scale=1000, stakeholder="the group controller",
         intro="After IntegrateCo's new accounting system, the controller studies how long the month-end close takes.",
         y=("Close duration", "working days"), x=("Journal entries", "thousand entries", "thousand entries", (2, 20)),
         d=("New ERP system", "months on the old system", "months on the new system"),
         b0=(3, 6), b1=(0.2, 0.6), b2=(-2.5, -0.5), b3=(-0.15, 0.05)),
    dict(title="Payroll overtime", phrase="{v} projects assigned", scale=1, stakeholder="Noah, IntegrateCo's CFO",
         intro="IntegrateCo (payroll case) wants to explain monthly overtime pay per employee.",
         y=("Overtime pay", "USD per month"), x=("Projects assigned", "projects", "project", (1, 8)),
         d=("Installation job", "employees in office roles", "employees in installation roles"),
         b0=(20, 120), b1=(30, 90), b2=(40, 200), b3=(-20, 40)),
    dict(title="P-card spending", phrase="{v} transactions in the month", scale=1, stakeholder="OSU's internal auditor",
         intro="OSU's internal audit team studies monthly P-card spending per cardholder (P-card case).",
         y=("Monthly P-card spending", "USD"), x=("Transactions", "transactions", "transaction", (2, 60)),
         d=("Student employee", "staff cardholders", "student-employee cardholders"),
         b0=(200, 900), b1=(150, 320), b2=(-900, -150), b3=(-120, 10)),
    dict(title="Journal entry review time", phrase="{v} line items in the entry", scale=1, stakeholder="the audit senior",
         intro="An audit team measures how long it takes to review journal entries for a client.",
         y=("Review time", "minutes per entry"), x=("Line items", "line items", "line item", (2, 30)),
         d=("Manual entry", "system-generated entries", "manual entries"),
         b0=(2, 6), b1=(0.5, 1.6), b2=(3, 12), b3=(-0.3, 0.8)),
    dict(title="Film permit lead time", phrase="{v} locations on the permit", scale=1, stakeholder="the NYC film office manager",
         intro="The NYC film office (Excel case) studies the lead time between entering and starting a permit.",
         y=("Lead time", "days"), x=("Locations on the permit", "locations", "location", (1, 12)),
         d=("Television shoot", "film shoots", "television shoots"),
         b0=(2, 5), b1=(0.3, 1.2), b2=(-2.5, -0.5), b3=(-0.4, 0.3)),
    dict(title="Invoice matching errors", phrase="{v} invoices processed", scale=100, stakeholder="the accounts payable manager",
         intro="An accounts payable team counts invoices that fail the three-way match (order, receipt, invoice).",
         y=("Matching errors", "errors per month"), x=("Invoices processed", "hundred invoices", "hundred invoices", (5, 60)),
         d=("Automated matching", "months with manual matching", "months with automated matching"),
         b0=(2, 8), b1=(0.6, 1.8), b2=(-6, -1), b3=(-0.9, 0.1)),
    dict(title="Accounts receivable collection", phrase="an invoice amount of {v} USD", scale=1000, stakeholder="the credit manager",
         intro="A wholesaler studies how many days customers take to pay their invoices.",
         y=("Days to pay", "days"), x=("Invoice amount", "thousand USD", "thousand USD of invoice value", (1, 50)),
         d=("New customer", "existing customers", "new customers"),
         b0=(20, 35), b1=(0.2, 0.9), b2=(3, 12), b3=(-0.3, 0.4)),
    dict(title="Music downloads", phrase="{v} marketing emails", scale=1, stakeholder="the Chinook sales manager",
         intro="The digital music store from the Chinook database studies monthly downloads per customer.",
         y=("Downloads", "tracks per month"), x=("Marketing emails", "emails", "marketing email", (0, 12)),
         d=("Premium member", "standard customers", "premium members"),
         b0=(3, 8), b1=(0.4, 1.5), b2=(2, 7), b3=(-0.4, 0.6)),
    dict(title="Audit hours", phrase="{v} subsidiaries", scale=1, stakeholder="the engagement partner",
         intro="An audit firm studies how many hours an engagement takes.",
         y=("Audit hours", "hours"), x=("Subsidiaries", "subsidiaries", "subsidiary", (1, 25)),
         d=("First-year client", "returning clients", "first-year clients"),
         b0=(150, 400), b1=(20, 60), b2=(40, 180), b3=(-8, 12)),
    dict(title="Audit fees", phrase="client revenue of {v} million DKK", scale=1, stakeholder="the audit partner",
         intro="An audit firm studies the fees it charges clients.",
         y=("Audit fee", "thousand DKK"), x=("Client revenue", "million DKK", "million DKK of client revenue", (10, 500)),
         d=("Listed client", "private clients", "listed clients"),
         b0=(50, 150), b1=(0.3, 1.2), b2=(60, 250), b3=(-0.3, 0.6)),
]


def scenario_exams(start=3, per_scenario=2, seed=2026):
    rng = random.Random(seed)
    exams, eid = [], start
    for s in SCENARIOS:
        for k in range(per_scenario):
            interaction = k % 2 == 1
            pick = lambda lo_hi, step: round(rng.uniform(*lo_hi) / step) * step
            span = s["b1"][1] - s["b1"][0]
            step1 = 0.01 if span < 1 else (0.1 if span < 10 else 1)
            coefs = [pick(s["b0"], 1 if s["b0"][1] > 10 else 0.1), pick(s["b1"], step1),
                     pick(s["b2"], 1 if abs(s["b2"][1]) > 10 or abs(s["b2"][0]) > 10 else 0.1)]
            pvals = [rng.choice([0.001, 0.004, 0.02, 0.31]), rng.choice([0.001, 0.003, 0.008, 0.02]),
                     rng.choice([0.002, 0.028, 0.041, 0.12, 0.27])]
            if interaction:
                coefs.append(pick(s["b3"], 0.01 if (s["b3"][1] - s["b3"][0]) < 2 else (0.1 if s["b3"][1] < 10 else 1)))
                pvals.append(rng.choice([0.006, 0.015, 0.09, 0.34]))
            lo, hi = s["x"][3]
            xv = round(rng.uniform(lo + (hi - lo) * 0.2, hi - (hi - lo) * 0.2))
            gx = round(rng.uniform(lo + (hi - lo) * 0.3, hi - (hi - lo) * 0.3))
            out = hi * 2 if hi > 2 else hi + 5
            exams.append(build_exam(
                f"REG{eid:02d}", s["title"],
                {"stakeholder": s["stakeholder"], "intro": s["intro"]},
                {"name": s["y"][0], "unit": s["y"][1]},
                {"name": s["x"][0], "unit": s["x"][1], "unit_one": s["x"][2], "range": s["x"][3],
                 "phrase": s["phrase"], "scale": s["scale"]},
                {"name": s["d"][0], "ref": s["d"][1], "comp": s["d"][2]},
                coefs, pvals, interaction, (xv, rng.choice([0, 1])), gx, (out, 0),
                "Constructed scenario for practice (coefficients are illustrative)."))
            eid += 1
    return exams


def all_exams():
    return restaurant_exams() + scenario_exams()
